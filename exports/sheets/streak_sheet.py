from openpyxl.styles import Font


class StreakSheet:

    def build(self, workbook, report):

        if report.streak is None:
            return

        ws = workbook.create_sheet("Streak")

        ws["A1"] = "Streak Analytics"
        ws["A1"].font = Font(bold=True, size=14)

        # --------------------------------------------------
        # Summary
        # --------------------------------------------------

        summary = [

            ("Longest Winning Streak", report.streak.longest_winning_streak),
            ("Longest Losing Streak", report.streak.longest_losing_streak),

            ("Average Winning Streak", report.streak.average_winning_streak),
            ("Average Losing Streak", report.streak.average_losing_streak),

            ("Total Winning Streaks", report.streak.total_winning_streaks),
            ("Total Losing Streaks", report.streak.total_losing_streaks),

            ("Current Streak Type", report.streak.current_streak_type),
            ("Current Streak Length", report.streak.current_streak_length),

            ("Largest Winning Streak P&L", report.streak.largest_winning_streak_pnl),
            ("Largest Losing Streak P&L", report.streak.largest_losing_streak_pnl),

        ]

        row = 3

        for metric, value in summary:

            ws.cell(row=row, column=1).value = metric
            ws.cell(row=row, column=2).value = value

            ws.cell(row=row, column=1).font = Font(bold=True)

            row += 1

        # --------------------------------------------------
        # Individual Streaks
        # --------------------------------------------------

        row += 2

        headers = [

            "Streak Type",
            "Length",
            "P&L"

        ]

        for col, header in enumerate(headers, start=1):

            ws.cell(row=row, column=col).value = header
            ws.cell(row=row, column=col).font = Font(bold=True)

        row += 1

        for streak in report.streak.streaks:

            ws.cell(row=row, column=1).value = streak.streak_type
            ws.cell(row=row, column=2).value = streak.length
            ws.cell(row=row, column=3).value = streak.pnl

            row += 1