from openpyxl.styles import Font


class CalendarSheet:

    def build(self, workbook, report):

        if report.calendar is None:
            return

        ws = workbook.create_sheet("Calendar")

        ws["A1"] = "Calendar Analytics"
        ws["A1"].font = Font(bold=True, size=14)

        headers = [
            "Trading Date",
            "Trades",
            "Wins",
            "Losses",
            "Win %",
            "Gross Profit",
            "Gross Loss",
            "Net Profit",
            "Avg P&L",
            "Best Trade",
            "Worst Trade",
            "Profit Factor",
            "Cumulative Equity"
        ]

        for col, header in enumerate(headers, start=1):
            ws.cell(row=3, column=col).value = header
            ws.cell(row=3, column=col).font = Font(bold=True)

        row = 4

        for day in report.calendar.days:

            ws.cell(row=row, column=1).value = day.trading_date
            ws.cell(row=row, column=2).value = day.total_trades
            ws.cell(row=row, column=3).value = day.winning_trades
            ws.cell(row=row, column=4).value = day.losing_trades
            ws.cell(row=row, column=5).value = day.win_rate
            ws.cell(row=row, column=6).value = day.gross_profit
            ws.cell(row=row, column=7).value = day.gross_loss
            ws.cell(row=row, column=8).value = day.net_profit
            ws.cell(row=row, column=9).value = day.average_pnl
            ws.cell(row=row, column=10).value = day.best_trade
            ws.cell(row=row, column=11).value = day.worst_trade
            ws.cell(row=row, column=12).value = day.profit_factor
            ws.cell(row=row, column=13).value = day.cumulative_equity

            row += 1