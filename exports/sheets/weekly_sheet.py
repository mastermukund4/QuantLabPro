from openpyxl.styles import Font


class WeeklySheet:

    def build(self, workbook, report):

        if report.weekly is None:
            return

        ws = workbook.create_sheet("Weekly")

        ws["A1"] = "Weekly Analytics"
        ws["A1"].font = Font(bold=True, size=14)

        headers = [
            "Week",
            "Trades",
            "Wins",
            "Losses",
            "Win %",
            "Gross Profit",
            "Gross Loss",
            "Net Profit",
            "Avg P&L",
            "Profit Factor"
        ]

        for col, header in enumerate(headers, start=1):
            ws.cell(row=3, column=col).value = header
            ws.cell(row=3, column=col).font = Font(bold=True)

        row = 4

        for week in report.weekly.weeks:

            ws.cell(row=row, column=1).value = week.week
            ws.cell(row=row, column=2).value = week.total_trades
            ws.cell(row=row, column=3).value = week.winning_trades
            ws.cell(row=row, column=4).value = week.losing_trades
            ws.cell(row=row, column=5).value = week.win_rate
            ws.cell(row=row, column=6).value = week.gross_profit
            ws.cell(row=row, column=7).value = week.gross_loss
            ws.cell(row=row, column=8).value = week.net_profit
            ws.cell(row=row, column=9).value = week.average_pnl
            ws.cell(row=row, column=10).value = week.profit_factor

            row += 1