from openpyxl.styles import Font


class MonthlySheet:

    def build(self, workbook, report):

        if report.monthly is None:
            return

        ws = workbook.create_sheet("Monthly")

        ws["A1"] = "Monthly Analytics"
        ws["A1"].font = Font(bold=True, size=14)

        headers = [
            "Month",
            "Trades",
            "Wins",
            "Losses",
            "Win %",
            "Gross Profit",
            "Gross Loss",
            "Net Profit",
            "Avg P&L",
            "Profit Factor"
        ]

        for col, header in enumerate(headers, start=1):
            ws.cell(row=3, column=col).value = header
            ws.cell(row=3, column=col).font = Font(bold=True)

        row = 4

        for month in report.monthly.months:

            ws.cell(row=row, column=1).value = month.month
            ws.cell(row=row, column=2).value = month.total_trades
            ws.cell(row=row, column=3).value = month.winning_trades
            ws.cell(row=row, column=4).value = month.losing_trades
            ws.cell(row=row, column=5).value = month.win_rate
            ws.cell(row=row, column=6).value = month.gross_profit
            ws.cell(row=row, column=7).value = month.gross_loss
            ws.cell(row=row, column=8).value = month.net_profit
            ws.cell(row=row, column=9).value = month.average_pnl
            ws.cell(row=row, column=10).value = month.profit_factor

            row += 1