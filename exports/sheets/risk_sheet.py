from openpyxl.styles import Font


class RiskSheet:

    def build(self, workbook, report):

        if report.risk is None:
            return

        ws = workbook.create_sheet("Risk")

        ws["A1"] = "Risk Analytics"
        ws["A1"].font = Font(bold=True, size=14)

        metrics = [

            ("Standard Deviation", report.risk.standard_deviation),
            ("Downside Deviation", report.risk.downside_deviation),
            ("Sharpe Ratio", report.risk.sharpe_ratio),
            ("Sortino Ratio", report.risk.sortino_ratio),
            ("Calmar Ratio", report.risk.calmar_ratio),
            ("Ulcer Index", report.risk.ulcer_index),
            ("Recovery Factor", report.risk.recovery_factor),
            ("Value at Risk (95%)", report.risk.value_at_risk_95),
            ("Conditional VaR (95%)", report.risk.conditional_var_95),
            ("Kelly %", report.risk.kelly_percent),
            ("Half Kelly %", report.risk.half_kelly_percent),

        ]

        row = 3

        for metric, value in metrics:

            ws.cell(row=row, column=1).value = metric
            ws.cell(row=row, column=2).value = value

            ws.cell(row=row, column=1).font = Font(bold=True)

            row += 1