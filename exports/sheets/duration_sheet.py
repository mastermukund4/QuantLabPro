from openpyxl.styles import Font


class DurationSheet:

    def build(self, workbook, report):

        if report.duration is None:
            return

        ws = workbook.create_sheet("Duration")

        ws["A1"] = "Duration Analytics"
        ws["A1"].font = Font(bold=True, size=14)

        headers = [
            "Duration Bucket",
            "Trades",
            "Wins",
            "Losses",
            "Win %",
            "Gross Profit",
            "Gross Loss",
            "Net Profit",
            "Avg P&L",
            "Best Trade",
            "Worst Trade",
            "Profit Factor"
        ]

        for col, header in enumerate(headers, start=1):
            ws.cell(row=3, column=col).value = header
            ws.cell(row=3, column=col).font = Font(bold=True)

        row = 4

        for duration in report.duration.durations:

            ws.cell(row=row, column=1).value = duration.bucket
            ws.cell(row=row, column=2).value = duration.total_trades
            ws.cell(row=row, column=3).value = duration.winning_trades
            ws.cell(row=row, column=4).value = duration.losing_trades
            ws.cell(row=row, column=5).value = duration.win_rate
            ws.cell(row=row, column=6).value = duration.gross_profit
            ws.cell(row=row, column=7).value = duration.gross_loss
            ws.cell(row=row, column=8).value = duration.net_profit
            ws.cell(row=row, column=9).value = duration.average_pnl
            ws.cell(row=row, column=10).value = duration.best_trade
            ws.cell(row=row, column=11).value = duration.worst_trade
            ws.cell(row=row, column=12).value = duration.profit_factor

            row += 1