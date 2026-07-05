from openpyxl.styles import Font


class PerformanceSheet:

    def build(self, workbook, report):

        ws = workbook.create_sheet("Performance")

        ws["A1"] = "PERFORMANCE"
        ws["A1"].font = Font(size=16, bold=True)

        rows = [

            ("Total Trades", report.activity.total_trades),

            ("Winning Trades", report.activity.winning_trades),

            ("Losing Trades", report.activity.losing_trades),

            ("Gross Profit", report.returns.gross_profit),

            ("Gross Loss", report.returns.gross_loss),

            ("Net Profit", report.returns.net_profit),

            ("Average Winning Trade", report.returns.average_winning_trade),

            ("Average Losing Trade", report.returns.average_losing_trade),

            ("Average P&L / Trade", report.returns.average_pnl_per_trade),

            ("Profit Factor", report.edge.profit_factor),

            ("Expectancy (R)", report.edge.expectancy_r),

        ]

        row = 3

        for title, value in rows:

            ws.cell(row=row, column=1).value = title

            ws.cell(row=row, column=2).value = value

            row += 1

        ws.column_dimensions["A"].width = 35
        ws.column_dimensions["B"].width = 20