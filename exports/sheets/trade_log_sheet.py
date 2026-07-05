from openpyxl.styles import Font


class TradeLogSheet:

    def build(self, workbook, trades):

        ws = workbook.create_sheet("Trade Log")

        headers = [
            "Trade ID",
            "Entry Date",
            "Entry Time",
            "Exit Date",
            "Exit Time",
            "Strategy",
            "Symbol",
            "Expiry Type",
            "P&L",
            "No. of Legs",
            "Cumulative P&L"
        ]

        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True)

        cumulative = 0

        for row, trade in enumerate(trades, start=2):

            cumulative += trade.total_pnl

            ws.cell(row=row, column=1).value = trade.trade_id
            ws.cell(row=row, column=2).value = str(trade.entry_date)
            ws.cell(row=row, column=3).value = str(trade.entry_time)

            ws.cell(row=row, column=4).value = str(trade.exit_date)
            ws.cell(row=row, column=5).value = str(trade.exit_time)

            ws.cell(row=row, column=6).value = trade.strategy
            ws.cell(row=row, column=7).value = trade.symbol
            ws.cell(row=row, column=8).value = trade.expiry_type

            ws.cell(row=row, column=9).value = trade.total_pnl
            ws.cell(row=row, column=10).value = len(trade.legs)
            ws.cell(row=row, column=11).value = cumulative

        widths = {
            "A": 12,
            "B": 15,
            "C": 12,
            "D": 15,
            "E": 12,
            "F": 25,
            "G": 15,
            "H": 15,
            "I": 15,
            "J": 10,
            "K": 18,
        }

        for col, width in widths.items():
            ws.column_dimensions[col].width = width