from openpyxl.styles import Font

from exports.styles.formatter import Formatter


class DashboardSheet:

    def build(self, workbook, report):

        fmt = Formatter()

        ws = workbook.create_sheet("Dashboard")

        # -------------------------------------------------
        # TITLE
        # -------------------------------------------------

        ws.merge_cells("A1:D1")

        ws["A1"] = "QUANTLAB PRO DASHBOARD"

        fmt.title(ws["A1"])

        # -------------------------------------------------
        # KPI SECTION
        # -------------------------------------------------

        rows = [

            ("Starting Capital", report.equity.starting_capital),

            ("Ending Capital", report.equity.ending_capital),

            ("Net Profit", report.returns.net_profit),

            ("Profit Factor", report.edge.profit_factor),

            ("Win Rate (%)", report.activity.win_rate),

            ("Max Drawdown (₹)", report.drawdown.max_drawdown_rupees),

            ("Max Drawdown (%)", report.drawdown.max_drawdown_percent),

            ("Expectancy (R)", report.edge.expectancy_r)

        ]

        start_row = 3

        for title, value in rows:

            ws.cell(row=start_row, column=1).value = title

            header = ws.cell(row=start_row, column=1)

            header.font = Font(bold=True)

            value_cell = ws.cell(row=start_row, column=2)

            value_cell.value = value

            start_row += 1

        # Currency rows

        for r in [3, 4, 5, 8]:

            fmt.currency(ws.cell(row=r, column=2))

        fmt.auto_width(ws)

        ws.freeze_panes = "A3"