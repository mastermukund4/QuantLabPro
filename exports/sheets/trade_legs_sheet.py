from openpyxl.styles import Font


class TradeLegsSheet:

    def build(self, workbook, trades):

        ws = workbook.create_sheet("Trade Legs")

        headers = [
            "Trade ID",
            "Leg No",
            "Buy / Sell",
            "Option Type",
            "Strike",
            "Quantity",
            "Entry Price",
            "Exit Price",
            "Leg P&L"
        ]

        for col, header in enumerate(headers, start=1):

            cell = ws.cell(row=1, column=col)
            cell.value = header
            cell.font = Font(bold=True)

        row = 2

        for trade in trades:

            for leg in trade.legs:

                ws.cell(row=row, column=1).value = trade.trade_id
                ws.cell(row=row, column=2).value = leg.leg_no
                ws.cell(row=row, column=3).value = leg.buy_sell
                ws.cell(row=row, column=4).value = leg.option_type
                ws.cell(row=row, column=5).value = leg.strike
                ws.cell(row=row, column=6).value = leg.quantity
                ws.cell(row=row, column=7).value = leg.entry_price
                ws.cell(row=row, column=8).value = leg.exit_price
                ws.cell(row=row, column=9).value = leg.pnl

                row += 1

        widths = {
            "A": 12,
            "B": 10,
            "C": 12,
            "D": 12,
            "E": 12,
            "F": 10,
            "G": 15,
            "H": 15,
            "I": 15
        }

        for col, width in widths.items():
            ws.column_dimensions[col].width = width