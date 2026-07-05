from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from exports.styles.colors import *


class Formatter:

    def title(self, cell):

        cell.font = Font(
            bold=True,
            size=18,
            color=TITLE_FONT
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=TITLE_FILL
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    def header(self, cell):

        cell.font = Font(
            bold=True,
            color=HEADER_FONT
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=HEADER_FILL
        )

    def currency(self, cell):

        cell.number_format = '"₹"#,##0.00'

    def percentage(self, cell):

        cell.number_format = '0.00%'

    def auto_width(self, worksheet):

        # Calculate maximum width for each column
        column_widths = {}

        for row in worksheet.iter_rows():

            for cell in row:

                # Ignore merged cells
                if cell.__class__.__name__ == "MergedCell":
                    continue

                if cell.value is None:
                    continue

                column_letter = get_column_letter(cell.column)

                length = len(str(cell.value))

                if column_letter not in column_widths:
                    column_widths[column_letter] = length
                else:
                    column_widths[column_letter] = max(
                        column_widths[column_letter],
                        length
                    )

        # Apply widths
        for column_letter, width in column_widths.items():

            worksheet.column_dimensions[column_letter].width = width + 3